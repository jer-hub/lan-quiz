"""Extract plain text from teacher-uploaded lesson files (LAN-only, no persistence)."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

MAX_UPLOAD_BYTES = 5 * 1024 * 1024  # 5 MB
MAX_EXTRACT_CHARS = 24_000  # ~6K tokens; headroom under Groq free TPM

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
ALLOWED_CONTENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
    "text/plain",
    "text/markdown",
    "text/x-markdown",
    "application/octet-stream",  # browsers sometimes omit type
}


@dataclass
class ExtractResult:
    text: str
    char_count: int
    truncated: bool
    warning: str
    filename: str


def _ext(filename: str) -> str:
    name = (filename or "").strip().lower()
    if "." not in name:
        return ""
    return "." + name.rsplit(".", 1)[-1]


def _normalize_whitespace(text: str) -> str:
    text = text.replace("\x00", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()


def _decode_plain(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        if page_text.strip():
            parts.append(page_text)
    return "\n\n".join(parts)


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for para in doc.paragraphs:
        t = (para.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [((c.text or "").strip()) for c in row.cells]
            cells = [c for c in cells if c]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(
    filename: str,
    content_type: str | None,
    data: bytes,
) -> ExtractResult:
    if not data:
        raise ValueError("Empty file")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError(f"File too large (max {MAX_UPLOAD_BYTES // (1024 * 1024)} MB)")

    ext = _ext(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file type. Use PDF, DOCX, TXT, or MD.")

    ctype = (content_type or "").split(";")[0].strip().lower()
    if ctype and ctype not in ALLOWED_CONTENT_TYPES and ext not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported content type")

    warning = ""
    if ext == ".pdf":
        raw = _extract_pdf(data)
        if len(raw.strip()) < 40:
            raise ValueError(
                "Could not extract enough text from this PDF. "
                "Scanned/image-only PDFs are not supported (no OCR). "
                "Try a text PDF, DOCX, or paste notes."
            )
    elif ext == ".docx":
        raw = _extract_docx(data)
        if not raw.strip():
            raise ValueError("DOCX contained no extractable text")
    elif ext in (".txt", ".md"):
        raw = _decode_plain(data)
    else:
        raise ValueError("Unsupported file type. Use PDF, DOCX, TXT, or MD.")

    text = _normalize_whitespace(raw)
    truncated = False
    if len(text) > MAX_EXTRACT_CHARS:
        text = text[:MAX_EXTRACT_CHARS].rsplit(" ", 1)[0] or text[:MAX_EXTRACT_CHARS]
        truncated = True
        warning = (
            f"Text truncated to {MAX_EXTRACT_CHARS:,} characters to fit free-tier token limits. "
            "Trim further or upload a shorter excerpt if needed."
        )

    if len(text) < 20:
        raise ValueError("Extracted text is too short to generate a quiz")

    return ExtractResult(
        text=text,
        char_count=len(text),
        truncated=truncated,
        warning=warning,
        filename=(filename or "upload")[:200],
    )
